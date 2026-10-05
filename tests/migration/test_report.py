import openpyxl

from migration import report
from migration import state as st
from migration.runner import Migrator
from migration.vimeo_source import build_video_plan
from tests.migration.fakes import FakeFetcher, FakeHost, FakeMux, FakeVimeo


def sheet_rows(path):
    wb = openpyxl.load_workbook(path)
    return wb, {ws.title: [[c.value for c in row] for row in ws.iter_rows()] for ws in wb.worksheets}


def test_report_after_migration_has_all_sheets_columns_and_summary(tmp_path):
    store, vimeo = st.StateStore("sqlite:///:memory:"), FakeVimeo()
    Migrator(store, FakeMux(), vimeo, FakeFetcher(tmp_path), FakeHost()).run(vimeo.items())
    path = report.write_workbook(str(tmp_path / "r.xlsx"), *report.records_from_store(store), "MIGRATION STATE")
    wb, rows = sheet_rows(path)
    assert wb.sheetnames == ["Videos", "Audio Tracks", "Caption Tracks", "Errors", "Summary"]
    assert rows["Videos"][0] == report.VIDEO_COLS
    assert rows["Audio Tracks"][0] == report.AUDIO_COLS and rows["Caption Tracks"][0] == report.CAPTION_COLS
    assert rows["Errors"][0] == report.ERROR_COLS
    video = dict(zip(report.VIDEO_COLS, rows["Videos"][1]))
    assert video["Migration Status"] == "COMPLETED" and video["Audio Track Count"] == 6 and video["Caption Track Count"] == 6
    assert video["Public Playback ID"] and video["Signed Playback ID"] and video["DRM Playback ID"]
    assert len(rows["Audio Tracks"]) == 7 and len(rows["Caption Tracks"]) == 7
    summary = dict(rows["Summary"][1:])
    assert summary["Videos completed"] == 1 and summary["Successfully migrated audio tracks"] == 6
    assert summary["Expected caption tracks"] == 6 and summary["Videos with all three playback IDs"] == 1


def test_dry_run_report_flags_missing_tracks_and_creates_no_mux_fields(tmp_path):
    api = FakeVimeo(audio_langs=("en", "fr"), caption_langs=tuple("en fr it pt ro es".split()))
    plans = [build_video_plan(api, api.videos["111"])]
    videos, tracks, errors = report.records_from_plans(plans)
    path = report.write_workbook(str(tmp_path / "d.xlsx"), videos, tracks, errors, "DRY RUN")
    _, rows = sheet_rows(path)
    v = dict(zip(report.VIDEO_COLS, rows["Videos"][1]))
    assert v["Migration Status"] == "PLANNED" and v["Mux Asset ID"] is None and v["Audio Track Count"] == 2
    statuses = [r[7] for r in rows["Audio Tracks"][1:]]
    assert statuses.count("MISSING_ON_VIMEO") == 4 and statuses.count("PLANNED") == 2
    assert len(rows["Errors"]) == 1 + 4
    summary = dict(rows["Summary"][1:])
    assert summary["Audio tracks found on Vimeo"] == 2 and summary["Failed audio tracks (failed or missing on Vimeo)"] == 4
    assert summary["Videos migrated (Mux asset created)"] == 0


def test_report_status_labels():
    base = {"status": "AUDIO_ATTACHING", "outcome": st.FAILED}
    assert report.report_status(base) == "FAILED (AUDIO_ATTACHING)"
    assert report.report_status({"status": "COMPLETED", "outcome": st.COMPLETED}) == "COMPLETED"
    assert report.report_status({"status": "ASSET_READY", "outcome": st.IN_PROGRESS}) == "ASSET_READY"


def test_workbook_is_cumulative_complete_and_refreshed_after_each_video(tmp_path, monkeypatch):
    from migration.folder1 import seed_state
    from tests.migration.fakes import video_json
    vimeo = FakeVimeo(videos={str(i): video_json(str(i)) for i in range(1, 6)})
    store = st.StateStore(f"sqlite:///{(tmp_path / 's.db').as_posix()}")
    path = str(tmp_path / report.LATEST_NAME)
    cp = report.Checkpointer(store, path)
    reasons = []
    inner = cp.__call__
    monkeypatch.setattr(cp.__class__, "__call__", lambda self, reason="x": (reasons.append(reason), inner(reason))[1])

    seed_state(vimeo, store, vimeo.items(), sniff=False)
    cp("seeded")
    _, rows = sheet_rows(path)
    assert len(rows["Videos"]) == 1 + 5 and len(rows["Audio Tracks"]) == 1 + 30 and len(rows["Caption Tracks"]) == 1 + 30

    m = Migrator(store, FakeMux(), vimeo, FakeFetcher(tmp_path), FakeHost(), checkpoint=cp)
    m.run(vimeo.items(), limit=2)
    assert any(r.endswith("COMPLETED") for r in reasons) and any("AUDIO_VERIFIED" in r for r in reasons)
    _, rows = sheet_rows(path)
    assert len(rows["Videos"]) == 1 + 5 and len(rows["Audio Tracks"]) == 1 + 30    # still all videos, not just migrated ones
    statuses = [dict(zip(report.VIDEO_COLS, r))["Migration Status"] for r in rows["Videos"][1:]]
    assert statuses.count("COMPLETED") == 2 and statuses.count("DISCOVERED") == 3
    summary = dict(rows["Summary"][1:])
    assert summary["Videos completed"] == 2 and summary["Successfully migrated audio tracks"] == 12
    assert "Last updated" in summary

    m.run(vimeo.items())     # finish the rest: the same file now holds everything
    _, rows = sheet_rows(path)
    assert dict(rows["Summary"][1:])["Videos completed"] == 5 and dict(rows["Summary"][1:])["Successfully migrated caption tracks"] == 30


def test_checkpoint_survives_a_locked_workbook(tmp_path, monkeypatch):
    store = st.StateStore("sqlite:///:memory:")
    cp = report.Checkpointer(store, str(tmp_path / "x.xlsx"))
    monkeypatch.setattr(report.os, "replace", lambda a, b: (_ for _ in ()).throw(PermissionError("locked")))
    assert cp("t") is None and not list(tmp_path.glob("*.writing.xlsx"))
